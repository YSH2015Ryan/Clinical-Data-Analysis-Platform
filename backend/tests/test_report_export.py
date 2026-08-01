"""T7.6 Word/Markdown 导出测试（docs/10 §3 派生管线；DoD：格式可用、章节完整）。

- html_to_markdown：10 章节标题/表格/免责声明完整；图表以占位标注；
- context_to_docx：同一 ReportContext → docx，章节标题/免责声明/图像齐全；
- generate_report 端到端：markdown/docx → READY，存储产物可读，幂等。
"""

import io

import pytest
from app.engines.reporting.context import DISCLAIMER, SECTION_TITLES, build_report_context
from app.engines.reporting.docx_export import context_to_docx
from app.engines.reporting.markdown_export import html_to_markdown
from app.engines.reporting.render import render_html
from app.models.entities import Report
from app.services.storage import storage_from_settings
from app.workers import db as worker_db
from app.workers.report_tasks import _embed_images, generate_report, key_of_report
from docx import Document
from sqlalchemy.ext.asyncio import AsyncSession

from tests.test_analysis_orchestration import _create_run, _fetch_run
from tests.test_analysis_orchestration import seeded as _seeded_impl
from tests.test_analysis_orchestration import worker_sync_db as _ws_impl

# 共享 fixture 再导出
seeded = _seeded_impl
worker_sync_db = _ws_impl

ZIP_MAGIC = b"PK"


def _payload() -> dict:
    return {
        "run": {
            "task_type": "classification",
            "status": "COMPLETED",
            "selected_models": ["logreg"],
            "created_at": "2026-08-01",
            "finished_at": "2026-08-01",
        },
        "dataset": {"name": "clf"},
        "target_summary": {"n_samples": 60, "n_features": 2, "positives": 30, "prevalence": 0.5},
        "validation": {"strategy": "stratified_kfold", "n_splits": 3, "seed": 42},
        "task_detection": None,
        "quality_report": {
            "status": "PASSED",
            "blockers": 0,
            "warnings": 0,
            "issues": [],
            "column_profiles": [
                {
                    "name": "age",
                    "inferred_type": "continuous",
                    "missing_rate": 0.0,
                    "unique": 60,
                }
            ],
        },
        "transform_log": [],
        "preprocess_steps": [{"type": "scale", "params": {}}],
        "input_hash": "0" * 64,
        "table": [
            {
                "model_key": "logreg",
                "status": "COMPLETED",
                "metrics": {"auc": {"mean": 0.8, "std": 0.01}},
                "duration_ms": 100,
            }
        ],
        "comparison": {
            "ranking": ["logreg"],
            "recommendation": {
                "top": [{"model_key": "logreg", "rank": 1, "reason": "主指标最高（auc=0.8000）"}],
                "not_recommended": [],
                "primary_metric": "auc",
                "formula": "score = 0.6×主指标归一化",
            },
            "charts": [
                {
                    "kind": "roc",
                    "title": "ROC 曲线",
                    "data": {
                        "series": [
                            {
                                "model_key": "logreg",
                                "fpr": [0, 0.1, 1],
                                "tpr": [0, 0.9, 1],
                                "auc": 0.8,
                            }
                        ]
                    },
                }
            ],
        },
        "software": {"python": "3.11"},
    }


def _context_with_images() -> dict:
    context = build_report_context(_payload())
    _embed_images(context)
    return context


class TestMarkdownExport:
    def test_sections_complete(self) -> None:
        md = html_to_markdown(render_html(_context_with_images()))
        for title in SECTION_TITLES.values():
            assert title in md
        assert DISCLAIMER in md
        assert "附录：可复现性信息" in md

    def test_table_and_image_placeholder(self) -> None:
        md = html_to_markdown(render_html(_context_with_images()))
        assert "| ---" in md  # 表格转为 Markdown 表
        assert "data:image/png" not in md  # base64 不内联（体积）
        assert "PNG 图像见 HTML/PDF 版报告" in md

    def test_minimal_html(self) -> None:
        md = html_to_markdown("<h1>标题</h1><p>段落</p><ul><li>甲</li><li>乙</li></ul>")
        assert "# 标题" in md
        assert "- 甲" in md and "- 乙" in md


class TestDocxExport:
    def test_sections_and_image(self) -> None:
        doc = Document(io.BytesIO(context_to_docx(_context_with_images())))
        text = "\n".join(p.text for p in doc.paragraphs)
        for title in SECTION_TITLES.values():
            assert title in text
        assert DISCLAIMER in text
        assert len(doc.inline_shapes) >= 1  # 图表 PNG 嵌入
        assert len(doc.tables) >= 2  # 变量概况 / 指标表 / 软件版本表

    def test_missing_data_marks_na(self) -> None:
        payload = _payload()
        payload["target_summary"] = None
        payload["quality_report"] = None
        context = build_report_context(payload)
        doc = Document(io.BytesIO(context_to_docx(context)))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "本次分析不适用" in text


@pytest.mark.usefixtures("worker_sync_db")
class TestGenerateExportFormats:
    @pytest.mark.parametrize("fmt", ["markdown", "docx"])
    async def test_export_ready_idempotent(
        self, db_session: AsyncSession, seeded, fmt: str
    ) -> None:
        user, version = seeded
        run = await _create_run(db_session, user, version, ["logreg"])
        from app.workers.analysis_tasks import dispatch_analysis_run

        dispatch_analysis_run(str(run.id), ["logreg"])
        done, _ = _fetch_run(run.id)
        assert done.status == "COMPLETED"

        report = Report(run_id=run.id, format=fmt, created_by=user.id)
        db_session.add(report)
        await db_session.commit()

        generate_report.apply_async(args=[str(report.id)])
        factory = worker_db.get_sync_session_factory()
        with factory() as db:
            saved = db.get(Report, report.id)
            assert saved is not None and saved.status == "READY"
            blob = storage_from_settings().get(key_of_report(saved))

        if fmt == "markdown":
            text = blob.decode("utf-8")
            for title in SECTION_TITLES.values():
                assert title in text
        else:
            assert blob.startswith(ZIP_MAGIC)
            doc = Document(io.BytesIO(blob))
            text = "\n".join(p.text for p in doc.paragraphs)
            for title in SECTION_TITLES.values():
                assert title in text

        # 幂等
        generate_report.apply_async(args=[str(report.id)])
        with factory() as db:
            again = db.get(Report, report.id)
            assert again is not None and again.status == "READY"
            assert again.storage_uri == saved.storage_uri
