"""ReportContext → Word（docs/10 §3：python-docx 模板填充；章节与 HTML 同源）。

图表 PNG 直接嵌入（image_b64 解码）；缺数据章节写「本次分析不适用」；
免责声明与可复现附录强制存在（CLAUDE.md §8）。
"""

import base64
import io
from typing import Any

from docx import Document
from docx.shared import Inches

from app.engines.reporting.context import NOT_APPLICABLE

_CHART_WIDTH = Inches(5.5)


def _add_charts(doc: Any, charts: list[dict[str, Any]]) -> None:
    for chart in charts:
        b64 = chart.get("image_b64")
        title = chart.get("title") or chart.get("kind") or "图表"
        if b64:
            doc.add_picture(io.BytesIO(base64.b64decode(b64)), width=_CHART_WIDTH)
            doc.add_paragraph(f"图：{title}")
        else:
            doc.add_paragraph(f"[图表] {title}（本次分析不适用）")


def _add_table(doc: Any, header: list[str], rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for i, h in enumerate(header):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v


def _metrics_cells(metrics: dict[str, Any] | None) -> list[str]:
    if not metrics:
        return ["-"]
    cells = []
    for name, val in metrics.items():
        if isinstance(val, dict) and "mean" in val:
            std = val.get("std")
            cells.append(f"{name}={val['mean']:.4f}" + (f"±{std:.4f}" if std is not None else ""))
        else:
            cells.append(f"{name}={val}")
    return cells


def _add_section(doc: Any, section: dict[str, Any]) -> None:
    doc.add_heading(section["title"], level=2)
    if not section.get("applicable"):
        doc.add_paragraph(NOT_APPLICABLE)
        return
    key, data = section["key"], section["data"]
    if key == "overview":
        for k, label in (
            ("n_samples", "样本量"),
            ("n_features", "特征数"),
            ("positives", "阳性样本数"),
            ("prevalence", "阳性率"),
            ("events", "事件数"),
            ("median_followup", "中位随访时间"),
            ("label_mean", "标签均值"),
        ):
            if data.get(k) is not None:
                doc.add_paragraph(f"{label}：{data[k]}")
        _add_table(
            doc,
            ["变量", "类型", "缺失率", "唯一值"],
            [
                [
                    str(c.get("name")),
                    str(c.get("inferred_type")),
                    str(c.get("missing_rate")),
                    str(c.get("unique")),
                ]
                for c in data.get("columns") or []
            ],
        )
    elif key == "quality":
        doc.add_paragraph(
            f"校验结论：{data.get('status')}"
            f"（阻断 {data.get('blockers')}，警告 {data.get('warnings')}）"
        )
        for issue in data.get("issues") or []:
            doc.add_paragraph(str(issue), style="List Bullet")
    elif key == "preprocessing":
        _add_table(
            doc,
            ["步骤", "参数"],
            [[str(s.get("type")), str(s.get("params"))] for s in data.get("steps") or []],
        )
    elif key == "methods":
        doc.add_paragraph(f"任务类型：{data.get('task_type')}")
        doc.add_paragraph(f"模型列表：{'、'.join(data.get('selected_models') or [])}")
        if data.get("validation"):
            doc.add_paragraph(f"交叉验证：{data['validation']}")
    elif key == "performance":
        _add_table(
            doc,
            ["模型", "状态", "指标"],
            [
                [
                    str(r.get("model_key")),
                    str(r.get("status")),
                    "；".join(_metrics_cells(r.get("metrics"))),
                ]
                for r in data.get("table") or []
            ],
        )
        _add_charts(doc, data.get("charts") or [])
    elif key == "explainability":
        _add_charts(doc, data.get("charts") or [])
    elif key == "comparison":
        doc.add_paragraph(f"排名：{' > '.join(data.get('ranking') or [])}")
        for entry in data.get("top") or []:
            doc.add_paragraph(
                f"第 {entry.get('rank')} 位 {entry.get('model_key')}：{entry.get('reason')}",
                style="List Bullet",
            )
        if data.get("formula"):
            doc.add_paragraph(f"排名公式：{data['formula']}")
    elif key == "clinical_interpretation":
        doc.add_paragraph(str(data.get("notice")))
        for stmt in data.get("statements") or []:
            doc.add_paragraph(str(stmt), style="List Bullet")
    elif key == "limitations":
        for item in data.get("items") or []:
            doc.add_paragraph(str(item), style="List Bullet")
    elif key == "conclusion":
        doc.add_paragraph(str(data.get("text")))


def context_to_docx(context: dict[str, Any]) -> bytes:
    """ReportContext（含 image_b64）→ docx 字节。"""
    doc = Document()
    dataset_name = (context.get("dataset") or {}).get("name") or "-"
    doc.add_heading(f"临床数据分析报告 — {dataset_name}", level=1)
    for section in context["sections"]:
        _add_section(doc, section)

    appendix = context["appendix"]
    doc.add_heading("附录：可复现性信息", level=2)
    if appendix.get("input_hash"):
        doc.add_paragraph(f"输入数据 hash：{appendix['input_hash']}")
    _add_table(
        doc,
        ["软件", "版本"],
        [[str(k), str(v)] for k, v in (appendix.get("software") or {}).items()],
    )
    doc.add_paragraph(str(appendix["disclaimer"]))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
